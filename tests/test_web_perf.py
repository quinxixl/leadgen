"""Performance changes of the web cabinet keep the same results: static caching, dashboard, lists, export."""
import csv
import io
import json
import os
import random
import unittest
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from unittest.mock import patch

from leadgen.core import TOPICS, stored_lead
from leadgen.web import (EXPORT_HEADINGS, STATIC_MAX_AGE, TTLCache, create_app, export_csv, export_xlsx,
                         list_lead, static_version)


def public_app(**config):
    return create_app({'TESTING':True,'SECRET_KEY':'test-'*10,'BOT_USERNAME':'test_bot','SESSION_COOKIE_SECURE':False}|config)


class StaticFiles(unittest.TestCase):
    def setUp(self):
        self.app=public_app()
        self.client=self.app.test_client()

    def test_templates_link_versioned_static_urls(self):
        page=self.client.get('/pricing').text
        version=static_version(self.app.static_folder,'app.css')
        self.assertEqual(len(version),10)
        self.assertIn(f'/web_static/app.css?v={version}',page)
        self.assertIn('/web_static/motion.js?v=',page)

    def test_versioned_url_is_cached_for_a_year(self):
        version=static_version(self.app.static_folder,'app.css')
        with self.client.get(f'/web_static/app.css?v={version}') as response:
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.cache_control.max_age,STATIC_MAX_AGE)
            self.assertTrue(response.cache_control.immutable)
            self.assertTrue(response.cache_control.public)

    def test_unversioned_or_stale_url_is_not_cached_long(self):
        for path in ('/web_static/app.css','/web_static/app.css?v=0123456789'):
            with self.client.get(path) as response:
                self.assertEqual(response.status_code,200,path)
                self.assertNotEqual(response.cache_control.max_age,STATIC_MAX_AGE,path)
                self.assertFalse(response.cache_control.immutable,path)
                self.assertNotIn('no-store',response.headers.get('Cache-Control',''),path)

    def test_static_skips_session_and_database(self):
        def no_database():raise AssertionError('static file opened a database connection')
        client=public_app(DB_FACTORY=no_database).test_client()
        with client.session_transaction() as s:s.update(user_id=1,csrf='x')
        with client.get('/web_static/app.css') as response:
            self.assertEqual(response.status_code,200)
            self.assertNotIn('Set-Cookie',response.headers)
            self.assertNotIn('Cookie',response.headers.get('Vary',''))
        # Pages still get the session cookie and no-store.
        page=self.app.test_client().get('/')
        self.assertIn('signalid_session=',page.headers['Set-Cookie'])
        self.assertEqual(page.headers['Cache-Control'],'no-store')

    def test_missing_file_and_traversal_have_no_version(self):
        self.assertEqual(static_version(self.app.static_folder,'missing.css'),'')
        self.assertEqual(static_version(self.app.static_folder,'../web.py'),'')
        self.assertEqual(static_version(self.app.static_folder,'tg'),'')
        self.assertEqual(self.client.get('/web_static/missing.css').status_code,404)


class Cache(unittest.TestCase):
    def test_entries_expire_and_size_is_bounded(self):
        now=[0.0];calls=[]
        cache=TTLCache(300,2,clock=lambda:now[0])
        load=lambda key:(lambda:calls.append(key) or key*10)
        self.assertEqual(cache.get(1,load(1)),10)
        self.assertEqual(cache.get(1,load(1)),10)
        self.assertEqual(calls,[1])
        now[0]=301.0
        self.assertEqual(cache.get(1,load(1)),10)
        self.assertEqual(calls,[1,1])
        cache.get(2,load(2));cache.get(3,load(3))
        self.assertEqual(len(cache._items),2)
        self.assertNotIn(1,cache._items)


def enriched_payload(i):
    return {'source':f'Telegram: chat{i%3}','url':f'https://t.me/c/{i}','title':f'Нужен сайт {i}' if i%5 else '=1+2',
            'text':'Нужен сайт для школы, бюджет 50 000 ₽ '*20,'published':'2026-10-01T10:00:00+00:00',
            'temperature':'горячий','score':70+i,'summary':'Сайт для школы','score_reason':'бюджет','budget_text':'50 000 ₽'}


def legacy_rows(payloads,statuses=('new','won','in_progress'),amounts=(None,15000,0)):
    return [{'id':i,'lead':stored_lead(p),'pipeline_status':statuses[i%3],'deal_amount':amounts[i%3]}
            for i,p in enumerate(payloads)]


def legacy_values(rows):
    from leadgen.product import PIPELINE
    def safe(v):
        v=str(v or '')
        return "'"+v if v.lstrip().startswith(('=','+','-','@','\t','\r')) else v
    return [[safe(x) for x in [r['lead'].title,r['lead'].source,r['lead'].url,r['lead'].published,r['lead'].score,
             PIPELINE[r['pipeline_status']],r['deal_amount']]] for r in rows]


def legacy_csv(rows):
    stream=io.StringIO();writer=csv.writer(stream,delimiter=';');writer.writerow(list(EXPORT_HEADINGS))
    writer.writerows(legacy_values(rows))
    return '﻿'+stream.getvalue()


def legacy_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='Лиды';sheet.append(list(EXPORT_HEADINGS))
    for values_row in legacy_values(rows):
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('',v)[:32767] for v in values_row])
    for cell in sheet[1]:
        cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='08182F')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for i,width in enumerate((55,32,45,28,12,24,20),1):sheet.column_dimensions[get_column_letter(i)].width=width
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:cell.alignment=Alignment(vertical='top',wrap_text=True)
    buffer=io.BytesIO();book.save(buffer)
    return buffer.getvalue()


def workbook_snapshot(data):
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(data))
    sheet=book.active
    cells=[[(c.value,c.data_type,c.font.b,c.font.color.rgb if c.font.color else None,c.fill.fgColor.rgb,
             c.alignment.vertical,c.alignment.wrap_text) for c in row] for row in sheet.iter_rows()]
    widths={key:dim.width for key,dim in sheet.column_dimensions.items() if dim.width}
    return book.sheetnames,cells,sheet.freeze_panes,sheet.auto_filter.ref,widths


class Export(unittest.TestCase):
    def setUp(self):
        payloads=[enriched_payload(i) for i in range(7)]
        payloads.append({'source':'FL','url':'https://fl.ru/1','title':'Бот\x01 для заявок','text':'Нужен чат-бот, бюджет 30000',
                         'published':None})  # stored without a score: enrich() computes it
        self.rows=legacy_rows(payloads)

    def test_csv_is_identical_and_streamed_in_chunks(self):
        self.assertEqual(''.join(export_csv(self.rows)),legacy_csv(self.rows))
        many=self.rows*200
        with patch('leadgen.web.CSV_CHUNK_ROWS',500):
            chunks=list(export_csv(many))
        self.assertEqual(len(chunks),4)
        self.assertEqual(''.join(chunks),legacy_csv(many))
        self.assertEqual(''.join(export_csv([])),legacy_csv([]))

    def test_xlsx_has_identical_cells_and_layout(self):
        for rows in (self.rows,[]):
            self.assertEqual(workbook_snapshot(export_xlsx(rows)),workbook_snapshot(legacy_xlsx(rows)))

    def test_list_lead_matches_full_payload(self):
        from leadgen.core import lead_data
        for payload in [enriched_payload(1),{'source':'FL','url':'u','title':'Нужен бот','text':'Нужен чат-бот, бюджет 30000','published':None}]:
            light=dict(payload)
            if payload.get('summary'):light.pop('text')
            full=stored_lead(payload)
            lead=list_lead(json.dumps(light))
            for field in ('title','source','url','published','score','summary','temperature'):
                self.assertEqual(getattr(lead,field),getattr(full,field),field)
            self.assertEqual(lead_data(json.dumps(light)),light)


class Workspaces(unittest.TestCase):
    """load_workspace writes only when the user's own workspace is missing."""
    class Db:
        def __init__(self,rows):self.rows=rows;self.queries=[];self.depth=0
        def execute(self,query,params=()):
            self.queries.append(query)
            if 'INSERT INTO workspaces' in query:self.rows=self.rows+[{'id':9,'name':'Моя','owner_user_id':1,'role':'owner'}]
            return SimpleNamespace(fetchall=lambda:self.rows,fetchone=lambda:{'id':9})
        def __enter__(self):return self
        def __exit__(self,*args):return False

    def test_existing_workspace_is_read_only(self):
        from leadgen.teams import load_workspace
        db=self.Db([{'id':5,'name':'Своя','owner_user_id':1,'role':'owner'},{'id':6,'name':'Чужая','owner_user_id':2,'role':'manager'}])
        rows,chosen=load_workspace(db,{'id':1,'display_name':'A'},'6')
        self.assertEqual(chosen['id'],6);self.assertEqual(len(db.queries),1)
        self.assertFalse(any('INSERT' in q for q in db.queries))

    def test_missing_own_workspace_is_created(self):
        from leadgen.teams import load_workspace
        db=self.Db([{'id':6,'name':'Чужая','owner_user_id':2,'role':'manager'}])
        rows,chosen=load_workspace(db,{'id':1,'display_name':'A'})
        self.assertTrue(any('INSERT INTO workspaces' in q for q in db.queries))
        self.assertIn(9,[row['id'] for row in rows])


class Webhooks(unittest.TestCase):
    def test_no_active_webhooks_skips_lead_snapshot(self):
        from leadgen.integrations import enqueue_lead_event
        queries=[]
        db=SimpleNamespace(execute=lambda q,p=():queries.append(q) or SimpleNamespace(fetchall=lambda:[],fetchone=lambda:None))
        with patch('leadgen.integrations.integrations_enabled',return_value=True):
            self.assertEqual(enqueue_lead_event(db,1,2),0)
        self.assertEqual(len(queries),1);self.assertIn('FROM webhooks',queries[0])


class PaymentRefresh(unittest.TestCase):
    def test_read_transaction_ends_before_api_call(self):
        from leadgen import payments
        events=[];row={'status':'pending','yookassa_payment_id':'yk-1','id':'x'}
        db=SimpleNamespace(execute=lambda q,p=():SimpleNamespace(fetchone=lambda:row),release=lambda:events.append('release'))
        client=SimpleNamespace(get_payment=lambda pid:events.append('api') or {'id':pid})
        with patch.object(payments,'process_payment',return_value='done'):
            payments.refresh(db,client,'6f1c1f0e-1e3c-4c9f-9d6a-1d2b3c4d5e6f',1)
        self.assertEqual(events,['release','api'])


class ReaderCounts(unittest.TestCase):
    def test_counts_come_from_the_given_function(self):
        from leadgen.telegram_monitor import reader_status
        db=SimpleNamespace(execute=lambda q,p=():SimpleNamespace(fetchone=lambda:None))
        status=reader_status(db,counts=lambda conn,now:(7,3))
        self.assertEqual((status['realtime_sources_24h'],status['joined_sources']),(7,3))


class Billing(unittest.TestCase):
    def test_poll_backs_off_from_2_to_5_seconds(self):
        from pathlib import Path
        script=Path('leadgen/web_static/billing.js').read_text()
        self.assertIn('POLL_FIRST=2000,POLL_STEP=1000,POLL_MAX=5000,POLL_FOR=60000',script)
        self.assertIn('delay=Math.min(POLL_MAX,delay+POLL_STEP)',script)


class Deploy(unittest.TestCase):
    def test_gunicorn_threads_match_pool_size(self):
        from pathlib import Path
        from leadgen.database import POOL_MAX_DEFAULT
        compose=Path('compose.web.yaml').read_text()
        self.assertIn('"--workers", "2", "--threads", "4"',compose)
        self.assertIn('DB_POOL_MAX: "4"',compose)
        self.assertGreaterEqual(POOL_MAX_DEFAULT,4)


# --- Postgres: dashboard equality, query counts, cached sources ----------------------------------------------
DSN=os.environ.get('WEB_TEST_DATABASE_URL','')

LEGACY_BASE="ul.user_id=? AND ul.delivery_status<>'filtered' AND ul.updated_at IS NOT NULL AND l.first_seen >= now()-(? * interval '1 day')"


def legacy_dashboard(db,owner_id,days):
    owner=(owner_id,days);base=LEGACY_BASE
    stats=db.execute("""SELECT count(*) AS total,
        count(*) FILTER (WHERE ul.pipeline_status='won') AS won,
        count(*) FILTER (WHERE f.label='fit') AS fit,
        sum(ul.deal_amount) FILTER (WHERE ul.pipeline_status='won') AS revenue
        FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
        LEFT JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
        WHERE """+base,owner).fetchone()
    stages=db.execute("SELECT ul.pipeline_status,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE "+base+" GROUP BY ul.pipeline_status",owner).fetchall()
    sources=db.execute("""SELECT l.source AS name,count(*) AS count,
        count(f.label) AS reviewed,count(*) FILTER (WHERE f.label='fit') AS fit,
        count(*) FILTER (WHERE f.label='ad') AS ads,
        count(*) FILTER (WHERE ul.pipeline_status='won') AS won
        FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
        LEFT JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
        WHERE """+base+" GROUP BY name ORDER BY won DESC,fit DESC,count DESC LIMIT 20",owner).fetchall()
    daily=db.execute("SELECT (l.first_seen AT TIME ZONE 'UTC')::date AS day,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE "+base+" GROUP BY day ORDER BY day",owner).fetchall()
    quality=db.execute("""SELECT f.label,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
        JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
        WHERE """+base+" GROUP BY f.label ORDER BY count DESC",owner).fetchall()
    response=db.execute("""SELECT avg(extract(epoch FROM (a.first_response-ul.sent_at)))/60 AS minutes,count(a.first_response) AS samples
        FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
        JOIN LATERAL (SELECT min(created_at) AS first_response FROM lead_activity
            WHERE user_id=ul.user_id AND lead_id=ul.lead_id AND kind='status' AND detail='Написали') a ON true
        WHERE """+base+" AND ul.sent_at IS NOT NULL AND a.first_response>=ul.sent_at",owner).fetchone()
    projects=db.execute("""SELECT p.name,count(*) AS count,count(*) FILTER (WHERE ul.pipeline_status='won') AS won
        FROM project_leads pl JOIN projects p ON p.id=pl.project_id AND p.user_id=pl.user_id
        JOIN user_leads ul ON ul.user_id=pl.user_id AND ul.lead_id=pl.lead_id JOIN leads l ON l.id=ul.lead_id
        WHERE """+base+" GROUP BY p.id ORDER BY count DESC",owner).fetchall()
    return {'stats':dict(stats),'stages':[dict(r) for r in stages],'sources':[dict(r) for r in sources],
            'daily':[dict(r) for r in daily],'quality':[dict(r) for r in quality],'response':dict(response),
            'projects':[dict(r) for r in projects]}


def as_set(rows):
    return sorted(tuple(sorted(row.items())) for row in rows)


if DSN:
    import test_web_postgres as web_postgres  # the module, not the class: its tests must not run twice here


@unittest.skipUnless(DSN,'WEB_TEST_DATABASE_URL is not configured')
class PerfPostgres(unittest.TestCase):
    if DSN:
        setUpClass=classmethod(web_postgres.WebPostgres.setUpClass.__func__)
        tearDownClass=classmethod(web_postgres.WebPostgres.tearDownClass.__func__)
        setUp=web_postgres.WebPostgres.setUp

    def seed(self,count=160,rng=None):
        rng=rng or random.Random(7)
        labels=['fit','ad','job','too_cold',None,None]
        statuses=['new','replied','in_progress','won','lost','archived']
        now=datetime.now(timezone.utc)
        with self.db:
            projects=[self.db.execute('''INSERT INTO projects(user_id,workspace_id,name,topics) VALUES(?,?,?,?::jsonb) RETURNING id''',
                                      (self.u,self.ws,name,json.dumps(['Сайты']))).fetchone()['id'] for name in ('Сайты','Боты','Пусто')]
            for i in range(count):
                seen=now-timedelta(days=rng.randint(0,100),hours=rng.randint(0,23))
                payload=enriched_payload(i)|{'source':f'Источник {rng.randint(0,25)}'}
                lid=self.db.execute('INSERT INTO leads(url,fingerprint,payload,status,reason,first_seen) VALUES(?,?,?,?,?,?) RETURNING id',
                    (payload['url']+'/s',f'seed-{i}',json.dumps(payload),'ready','запрос',seen)).fetchone()['id']
                owner=self.u if i%7 else self.other
                delivery=rng.choice(['sent','sent','sent','filtered','pending'])
                sent=seen+timedelta(minutes=rng.randint(1,90)) if delivery=='sent' else None
                status=rng.choice(statuses)
                self.db.execute('''INSERT INTO user_leads(user_id,lead_id,delivery_status,sent_at,deal_amount) VALUES(?,?,?,?,?)''',
                    (owner,lid,delivery,sent,rng.choice([None,1000,25000,0])))
                self.db.execute('UPDATE user_leads SET pipeline_status=? WHERE user_id=? AND lead_id=?',(status,owner,lid))
                label=rng.choice(labels)
                if label:self.db.execute('INSERT INTO lead_feedback(user_id,lead_id,label) VALUES(?,?,?)',(owner,lid,label))
                if sent and rng.random()<0.5:
                    for _ in range(rng.randint(1,2)):
                        self.db.execute("INSERT INTO lead_activity(user_id,lead_id,kind,detail,created_at) VALUES(?,?,'status','Написали',?)",
                                        (owner,lid,sent+timedelta(minutes=rng.randint(-30,600))))
                if owner==self.u and rng.random()<0.6:
                    for project in rng.sample(projects[:2],rng.randint(1,2)):
                        self.db.execute('INSERT INTO project_leads(project_id,user_id,lead_id) VALUES(?,?,?)',(project,self.u,lid))

    def test_dashboard_report_equals_legacy_queries(self):
        from leadgen.dashboard import dashboard_report
        self.seed()
        for owner in (self.u,self.other):
            for days in (7,30,90):
                new,old=dashboard_report(self.db,owner,days),legacy_dashboard(self.db,owner,days)
                label=(owner,days)
                self.assertGreater(old['stats']['total'],0,label)
                for key in ('stats','response'):self.assertEqual(new[key],old[key],(key,label))
                self.assertEqual(new['daily'],old['daily'],label)
                for key in ('stages','quality','projects'):self.assertEqual(as_set(new[key]),as_set(old[key]),(key,label))
                # Sources: same rows; ties in (won, fit, count) may come in either order, as before.
                rank=lambda row:(row['won'],row['fit'],row['count'])
                self.assertEqual(sorted(map(rank,new['sources']),reverse=True),[rank(r) for r in new['sources']],label)
                self.assertEqual(sorted(map(rank,new['sources'])),sorted(map(rank,old['sources'])),label)
                self.assertLessEqual(len(new['sources']),20)
                self.assertEqual(sorted(r['count'] for r in new['quality'])[::-1],[r['count'] for r in new['quality']])
                self.assertEqual(sorted(r['count'] for r in new['projects'])[::-1],[r['count'] for r in new['projects']])

    def test_dashboard_report_without_leads(self):
        from leadgen.dashboard import dashboard_report
        with self.db:self.db.execute('DELETE FROM user_leads WHERE user_id=?',(self.u,))
        report=dashboard_report(self.db,self.u,30)
        self.assertEqual(report,legacy_dashboard(self.db,self.u,30))
        self.assertEqual(self.client.get('/app').status_code,200)

    def counting_client(self):
        from leadgen.database import db_open
        from test_web_postgres import DSN as dsn
        queries=[]
        class Counting:
            def __init__(self):self.inner=db_open(dsn)
            def execute(self,query,params=()):queries.append(query);return self.inner.execute(query,params)
            def __getattr__(self,name):return getattr(self.inner,name)
            def __enter__(self):self.inner.__enter__();return self
            def __exit__(self,*args):return self.inner.__exit__(*args)
        app=create_app(dict(self.app.config)|{'DB_FACTORY':Counting})
        client=app.test_client()
        with client.session_transaction() as s:s.update(user_id=self.u,csrf='test-csrf')
        return app,client,queries

    def test_dashboard_page_scans_leads_once_and_renders_numbers(self):
        self.seed(60)
        app,client,queries=self.counting_client()
        r=client.get('/app?days=90')
        self.assertEqual(r.status_code,200)
        self.assertEqual(sum("first_seen >= now()-(" in q for q in queries),1)
        old=legacy_dashboard(self.db,self.u,90)
        self.assertIn(f"<strong>{old['stats']['total']}</strong>",r.text)
        for row in old['daily']:self.assertIn(row['day'].strftime('%d.%m.%Y'),r.text)
        # The workspace exists: no INSERT on a plain page view, and the session is not rewritten.
        self.assertFalse(any(q.lstrip().upper().startswith('INSERT') for q in queries))
        second=client.get('/app')
        self.assertNotIn('Set-Cookie',second.headers)

    def test_leads_list_caches_sources_and_leaves_text_out(self):
        app,client,queries=self.counting_client()
        first=client.get('/app/leads')
        self.assertEqual(first.status_code,200);self.assertIn('Telegram: test',first.text)
        distinct=lambda:sum('SELECT DISTINCT l.source' in q for q in queries)
        self.assertEqual(distinct(),1)
        client.get('/app/leads?page=1');self.assertEqual(distinct(),1)
        app.extensions['lead_source_names'].clear()
        client.get('/app/leads');self.assertEqual(distinct(),2)
        listed=[q for q in queries if 'ORDER BY l.id DESC LIMIT' in q]
        self.assertTrue(listed and all("l.payload-'text'" in q for q in listed))

    def test_export_matches_legacy_on_database_rows(self):
        self.seed(40)
        rows=self.db.execute('''SELECT l.id,l.payload,ul.pipeline_status,ul.deal_amount FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
            WHERE ul.user_id=? AND ul.delivery_status<>'filtered' ORDER BY l.id DESC''',(self.u,)).fetchall()
        legacy=[dict(r)|{'lead':stored_lead(r['payload'])} for r in rows]
        r=self.client.get('/app/leads/export')
        self.assertEqual(r.status_code,200);self.assertEqual(r.text,legacy_csv(legacy))
        x=self.client.get('/app/leads/export?format=xlsx')
        self.assertEqual(workbook_snapshot(x.data),workbook_snapshot(legacy_xlsx(legacy)))

    def test_admin_reader_counts_are_cached(self):
        app,client,queries=self.counting_client()
        with patch.dict(os.environ,{'WEB_ADMIN_TELEGRAM_IDS':'1'}):
            for _ in range(2):self.assertEqual(client.get('/admin').status_code,200)
        self.assertEqual(sum('selection_group=' in q for q in queries),1)

    def test_webhook_event_still_queued_for_active_hook(self):
        from leadgen.integrations import enqueue_lead_event
        with self.db,patch('leadgen.integrations.integrations_enabled',return_value=True):
            self.db.execute('''INSERT INTO webhooks(workspace_id,name,url,secret_cipher,created_by) VALUES(?,?,?,?,?)''',
                            (self.ws,'CRM','https://example.com/hook','x',self.u))
            self.assertEqual(enqueue_lead_event(self.db,self.u,self.ids[0],'lead.updated'),1)
            self.assertEqual(enqueue_lead_event(self.db,self.other,self.ids[1],'lead.updated'),0)
        row=self.db.execute('SELECT payload FROM webhook_outbox ORDER BY id DESC LIMIT 1').fetchone()
        self.assertEqual(row['payload']['lead']['id'],self.ids[0])


if __name__=='__main__':
    unittest.main()
