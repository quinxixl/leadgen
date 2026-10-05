"""Opt-in integration tests against a disposable local PostgreSQL database only."""
import json
import os
import unittest
from unittest.mock import patch
from pathlib import Path
from datetime import datetime,timezone
from leadgen.database import db_open
from leadgen.web import create_app
from leadgen.web_auth import approve_login
from cryptography.fernet import Fernet


class FakeTelegramGateway:
    def begin(self,phone):
        return {'phone':phone,'phone_code_hash':'hash','session':'pending-session'}
    def verify_code(self,state,code):
        if code=='2222':return {'password_required':True,'state':state|{'session':'password-session'}}
        return {'password_required':False,'session':'active-session','telegram_user_id':9001,'display_name':'Тестовый аккаунт'}
    def verify_password(self,state,password):
        if password!='correct':
            from leadgen.telegram_accounts import TelegramAccountError
            raise TelegramAccountError('Неверный облачный пароль Telegram.')
        return {'password_required':False,'session':'active-session','telegram_user_id':9001,'display_name':'Тестовый аккаунт'}
    def dialogs(self,session):
        return [{'peer_id':-10001,'title':'Бизнес-чат','username':'business_chat','kind':'supergroup'},
                {'peer_id':-10002,'title':'Автоматизации','username':None,'kind':'supergroup'}]
    def logout(self,session):return True


def fake_ai_generator(lead,profile_services='',portfolio='',settings=None,api_key=None,model=None,opener=None,user_id=None):
    return ({'Короткий':'Готов обсудить вашу задачу. Когда нужен результат?',
             'Экспертный':'Вижу проблему с потерей заявок. Предлагаю сначала описать текущий процесс.',
             'Дружелюбный':'Здравствуйте! Давайте разберём задачу и выберем подходящий вариант.'},
            {'prompt_tokens':321,'completion_tokens':123})

DSN=os.environ.get('WEB_TEST_DATABASE_URL','')


@unittest.skipUnless(DSN,'WEB_TEST_DATABASE_URL is not configured')
class WebPostgres(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from urllib.parse import urlsplit
        url=urlsplit(DSN)
        if url.hostname!='127.0.0.1' or url.port!=55439 or url.path!='/leadfinder_web_test':
            raise RuntimeError('Use only the disposable local test database on port 55439')
        cls.db=db_open(DSN)
        for role in ('anon','authenticated'):
            if not cls.db.execute('SELECT 1 FROM pg_roles WHERE rolname=?',(role,)).fetchone():
                cls.db.execute('CREATE ROLE '+role)
        for filename in ('20260929134042_create_leadgen_schema.sql','20260930092451_add_multiuser_product.sql'):
            cls.db.connection.execute((Path('supabase/migrations')/filename).read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.projects') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20260930143416_web_projects_crm.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.lead_reminders') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20260930193638_lead_reminders.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.telegram_connections') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261001041250_personal_telegram_connections.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.workspaces') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261001061514_team_workspaces_and_assignments.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.lead_tags') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261001062736_lead_feedback_and_tags.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.api_tokens') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261001063156_api_and_webhook_integrations.sql').read_text(),prepare=False)
        columns=cls.db.execute("SELECT 1 FROM information_schema.columns WHERE table_schema='leadgen' AND table_name='lead_assignments' AND column_name='notification_status'").fetchone()
        if not columns:cls.db.connection.execute(Path('supabase/migrations/20261001064130_assignment_notifications.sql').read_text(),prepare=False)
        columns=cls.db.execute("SELECT 1 FROM information_schema.columns WHERE table_schema='leadgen' AND table_name='projects' AND column_name='reply_sender'").fetchone()
        if not columns:cls.db.connection.execute(Path('supabase/migrations/20261001064937_project_reply_profiles.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.ai_offer_generations') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261001094118_ai_offer_generations.sql').read_text(),prepare=False)
        cls.db.connection.execute(Path('supabase/migrations/20261003120000_subscription_trials.sql').read_text(),prepare=False)
        if not cls.db.execute("SELECT to_regclass('leadgen.telegram_replies') AS name").fetchone()['name']:
            cls.db.connection.execute(Path('supabase/migrations/20261005070330_telegram_replies.sql').read_text(),prepare=False)
        cls.db.commit()
        cls.app=create_app({'TESTING':True,'SECRET_KEY':'test-'*10,'BOT_USERNAME':'test_bot',
            'SESSION_COOKIE_SECURE':False,'DB_FACTORY':lambda:db_open(DSN),
            'TELEGRAM_CIPHER_KEY':Fernet.generate_key().decode(),
            'TELEGRAM_GATEWAY_FACTORY':FakeTelegramGateway,
            'AI_OFFER_GENERATOR':fake_ai_generator,
            'WEBHOOK_RESOLVER':lambda *args,**kwargs:[(2,1,6,'',('93.184.216.34',443))]})

    @classmethod
    def tearDownClass(cls):cls.db.close()

    def setUp(self):
        with self.db:
            self.db.execute('TRUNCATE app_users CASCADE')
            self.db.execute('DELETE FROM leads')
            self.db.execute('DELETE FROM settings WHERE key LIKE ?',('web_login:%',))
            self.u=self.db.execute("INSERT INTO app_users(telegram_user_id,telegram_chat_id,display_name) VALUES(1,1,'Первый') RETURNING id").fetchone()['id']
            self.other=self.db.execute("INSERT INTO app_users(telegram_user_id,telegram_chat_id,display_name) VALUES(2,2,'Второй') RETURNING id").fetchone()['id']
            for uid in (self.u,self.other):
                self.db.execute('INSERT INTO user_preferences(user_id) VALUES(?)',(uid,))
                self.db.execute('INSERT INTO subscriptions(user_id) VALUES(?)',(uid,))
                workspace=self.db.execute("INSERT INTO workspaces(owner_user_id,name) VALUES(?,?) RETURNING id",
                                          (uid,'Команда '+str(uid))).fetchone()['id']
                self.db.execute("INSERT INTO workspace_members(workspace_id,user_id,role) VALUES(?,?,'owner')",(workspace,uid))
                if uid==self.u:self.ws=workspace
                else:self.other_ws=workspace
            self.ids=[]
            for i,uid in enumerate((self.u,self.other)):
                payload=dict(source='Telegram: test',url=f'https://t.me/test/{i}',title='Нужен сайт '+str(i),
                    text='<script>alert(1)</script> Нужен сайт',published=datetime.now(timezone.utc).isoformat())
                lid=self.db.execute('INSERT INTO leads(url,fingerprint,payload,status,reason,first_seen) VALUES(?,?,?,?,?,?) RETURNING id',
                    (payload['url'],str(i),json.dumps(payload),'ready','запрос',payload['published'])).fetchone()['id']
                self.db.execute("INSERT INTO user_leads(user_id,lead_id,delivery_status) VALUES(?,?,'sent')",(uid,lid));self.ids.append(lid)
        self.client=self.app.test_client()
        with self.client.session_transaction() as s:s.update(user_id=self.u,csrf='test-csrf')

    def test_render_owner_pages_and_hide_other_user(self):
        for path in ('/app','/app/leads','/app/settings','/app/sources',f'/app/leads/{self.ids[0]}'):
            r=self.client.get(path);self.assertEqual(r.status_code,200,path)
            self.assertNotIn('<script>alert(1)</script>',r.text)
        self.assertNotIn('Нужен сайт 1',self.client.get('/app/leads').text)
        self.assertEqual(self.client.get(f'/app/leads/{self.ids[1]}').status_code,404)
        self.assertEqual(self.client.get('/admin').status_code,403)

    def test_trial_then_expired_subscription_redirects_to_billing(self):
        sub=self.db.execute('SELECT status,ends_at FROM subscriptions WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual(sub['status'],'trial');self.assertIsNotNone(sub['ends_at'])
        self.assertEqual(self.client.get('/app/leads').status_code,200)
        with self.db:self.db.execute("UPDATE subscriptions SET ends_at=now()-interval '1 minute' WHERE user_id=?",(self.u,))
        r=self.client.get('/app/leads');self.assertEqual(r.status_code,302);self.assertTrue(r.location.endswith('/app/billing'))
        r=self.client.post(f'/app/leads/{self.ids[0]}',data={'csrf':'test-csrf','status':'won'})
        self.assertEqual(r.status_code,302)
        self.assertEqual(self.db.execute('SELECT pipeline_status FROM user_leads WHERE lead_id=?',(self.ids[0],)).fetchone()['pipeline_status'],'saved')
        self.assertEqual(self.client.get('/app/billing').status_code,200)

    def test_login_rate_limited_per_ip(self):
        with self.db:self.db.execute("DELETE FROM settings WHERE key LIKE ?",('rl:%',))
        c=self.app.test_client();c.get('/login')
        with c.session_transaction() as s:csrf=s['csrf']
        codes=[c.post('/login',data={'csrf':csrf,'action':'begin','consent':'1'}).status_code for _ in range(11)]
        self.assertEqual(codes[:10],[200]*10);self.assertEqual(codes[10],429)
        with self.db:self.db.execute("DELETE FROM settings WHERE key LIKE ?",('rl:%',))

    def test_cross_user_mutation_is_denied(self):
        r=self.client.post(f'/app/leads/{self.ids[1]}',data={'csrf':'test-csrf','status':'won','amount':'100'})
        self.assertEqual(r.status_code,404)
        self.assertEqual(self.db.execute('SELECT pipeline_status FROM user_leads WHERE lead_id=?',(self.ids[1],)).fetchone()['pipeline_status'],'saved')

    def test_status_feedback_amount_and_settings_persist(self):
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}',data={'csrf':'test-csrf','status':'won','amount':'75000','feedback':'fit'}).status_code,302)
        row=self.db.execute('SELECT * FROM user_leads WHERE lead_id=?',(self.ids[0],)).fetchone()
        self.assertEqual(row['deal_amount'],75000);self.assertEqual(row['pipeline_status'],'won')
        self.assertIn('75 000',self.client.get('/app').text)
        r=self.client.post('/app/settings',data={'csrf':'test-csrf','min_budget':'8000','topics':['Сайты'],'active':'on'})
        self.assertEqual(r.status_code,302)
        prefs=self.db.execute('SELECT * FROM user_preferences WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual(prefs['min_budget'],8000);self.assertEqual(prefs['topics'],['Сайты'])

    def test_login_flow(self):
        c=self.app.test_client();c.get('/login')
        with c.session_transaction() as s:csrf=s['csrf']
        self.assertEqual(c.post('/login',data={'csrf':csrf,'action':'begin','consent':'1'}).status_code,200)
        with c.session_transaction() as s:token,code=s['login_token'],s['login_code']
        approve_login(self.db,token,self.u,code)
        self.assertEqual(c.post('/login',data={'csrf':csrf,'action':'finish'}).status_code,302)
        self.assertEqual(c.get('/app').status_code,200)

    def test_export_is_scoped(self):
        r=self.client.get('/app/leads/export')
        self.assertEqual(r.status_code,200);self.assertNotIn('Нужен сайт 1',r.text)

    def test_projects_are_owner_scoped_and_notes_persist(self):
        r=self.client.post('/app/projects',data={'csrf':'test-csrf','name':'Разработка'})
        self.assertEqual(r.status_code,302)
        path=r.headers['Location']
        self.assertEqual(self.client.get(path).status_code,200)
        self.assertEqual(self.client.post(path,data={'csrf':'test-csrf','name':'Сайты','min_budget':'5000','min_score':'20',
            'topics':['Сайты'],'keywords':'сайт\nлендинг','enabled':'on','reply_sender':'Студия Север',
            'reply_offer':'Разрабатываем сайты','reply_proof':'Сайт для школы','reply_question':'Когда нужен запуск?',
            'reply_signature':'Анна','forbidden_phrases':'гарантия','reply_max_length':'700'}).status_code,302)
        project=self.db.execute('SELECT * FROM projects WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual(project['reply_sender'],'Студия Север')
        self.assertEqual(project['reply_proof'],['Сайт для школы'])
        with self.db:self.db.execute('INSERT INTO project_leads(project_id,user_id,lead_id) VALUES(?,?,?)',
                                     (project['id'],self.u,self.ids[0]))
        detail=self.client.get(f'/app/leads/{self.ids[0]}')
        self.assertIn('Студия Север',detail.text);self.assertIn('Когда нужен запуск?',detail.text)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}/notes',data={'csrf':'test-csrf','note':'Созвон в четверг'}).status_code,302)
        self.assertIn('Созвон в четверг',self.client.get(f'/app/leads/{self.ids[0]}').text)
        with self.client.session_transaction() as s:s['user_id']=self.other
        self.assertEqual(self.client.get(path).status_code,404)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}/notes',data={'csrf':'test-csrf','note':'Чужое'}).status_code,404)

    def test_two_projects_deliver_one_notification(self):
        from leadgen.product import deliver_registered
        with self.db:
            self.db.execute('DELETE FROM user_leads WHERE user_id=?',(self.u,))
            self.db.execute('UPDATE user_preferences SET monitoring_active=true WHERE user_id=?',(self.u,))
            self.db.execute("UPDATE leads SET first_seen=(now()+interval '1 second')::text")
            for name in ('Первый проект','Второй проект'):
                self.db.execute('''INSERT INTO projects(user_id,workspace_id,name,topics,show_without_budget)
                    VALUES(?,?,?,?::jsonb,true)''',(self.u,self.ws,name,json.dumps(['Сайты'])))
        calls=[]
        delivered=deliver_registered(self.db,{'sources':[],'max_age_hours':72,'min_budget':5000},lambda *args:calls.append(args))
        self.assertEqual(delivered,1);self.assertEqual(len(calls),1)
        count=self.db.execute('SELECT count(*) AS n FROM project_leads WHERE user_id=?',(self.u,)).fetchone()['n']
        self.assertEqual(count,2)

    def test_export_excel_and_filters_are_owner_scoped(self):
        import io
        from openpyxl import load_workbook
        lid=self.ids[0]
        with self.db:
            self.db.execute("INSERT INTO lead_feedback(user_id,lead_id,label) VALUES(?,?,'fit')",(self.u,lid))
            payload=json.loads(self.db.execute('SELECT payload FROM leads WHERE id=?',(lid,)).fetchone()['payload'])
            payload['title']='=HYPERLINK("https://example.com")'
            self.db.execute('UPDATE leads SET payload=? WHERE id=?',(json.dumps(payload),lid))
        r=self.client.get('/app/leads/export?format=xlsx&feedback=fit&source=Telegram:+test')
        self.assertEqual(r.status_code,200)
        sheet=load_workbook(io.BytesIO(r.data)).active
        self.assertEqual(sheet.max_row,2)
        self.assertNotEqual(sheet['A2'].data_type,'f')
        self.assertTrue(sheet['A2'].value.startswith("'="))
        self.assertEqual(sheet.freeze_panes,'A2')
        r=self.client.get('/app/leads/export?format=xlsx&feedback=ad')
        self.assertEqual(load_workbook(io.BytesIO(r.data)).active.max_row,1)
        self.assertEqual(self.client.get('/app/leads?from=bad').status_code,400)
        self.assertEqual(self.client.get('/app/leads?from=2026-09-30&to=2026-09-01').status_code,400)
        self.assertEqual(self.client.get('/app/leads?source=missing').status_code,200)
        self.assertEqual(self.client.get('/app/leads/export?from=2099-01-01').text.count('Нужен сайт'),0)

    def test_reminder_delivery_retry_and_ownership(self):
        from datetime import timedelta
        from leadgen.reminders import schedule,cancel,deliver_due
        future=datetime.now(timezone.utc)+timedelta(hours=1)
        lid=self.ids[0]
        with self.assertRaises(ValueError):schedule(self.db,self.u,self.ids[1],future)
        with self.assertRaises(ValueError):schedule(self.db,self.u,lid,future-timedelta(days=1))
        schedule(self.db,self.u,lid,future,'Уточнить задачу')
        calls=[]
        self.assertEqual(deliver_due(self.db,lambda *args:calls.append(args)),0)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[1]}/reminder',data={'csrf':'test-csrf','action':'cancel'}).status_code,404)
        with self.db:self.db.execute("UPDATE lead_reminders SET due_at=now()-interval '1 second' WHERE user_id=?",(self.u,))
        self.assertEqual(deliver_due(self.db,lambda *args:calls.append(args)),1)
        self.assertEqual(deliver_due(self.db,lambda *args:calls.append(args)),0)
        self.assertEqual(len(calls),1)
        self.assertEqual(str(calls[0][1]['chat_id']),'1')
        self.assertIn('Уточнить задачу',calls[0][1]['text'])
        schedule(self.db,self.u,lid,future)
        with self.db:self.db.execute("UPDATE lead_reminders SET due_at=now()-interval '1 second' WHERE user_id=?",(self.u,))
        def fail(*args):raise RuntimeError('network')
        self.assertEqual(deliver_due(self.db,fail),0)
        self.assertEqual(self.db.execute('SELECT status FROM lead_reminders WHERE user_id=?',(self.u,)).fetchone()['status'],'uncertain')
        self.assertEqual(deliver_due(self.db,lambda *args:calls.append(args)),0)
        self.assertIn('доставка не подтверждена',self.client.get(f'/app/leads/{lid}').text)
        self.assertTrue(cancel(self.db,self.u,lid))
        self.assertFalse(cancel(self.db,self.other,lid))
        response=self.client.post(f'/app/leads/{lid}/reminder',data={'csrf':'test-csrf','due':future.strftime('%Y-%m-%dT%H:%M')})
        self.assertEqual(response.status_code,302)

    def test_analytics_periods_and_revenue_are_scoped(self):
        with self.db:
            self.db.execute("UPDATE user_leads SET pipeline_status='won',deal_amount=77001 WHERE user_id=?",(self.u,))
            self.db.execute("UPDATE user_leads SET pipeline_status='won',deal_amount=998877 WHERE user_id=?",(self.other,))
        page=self.client.get('/app?days=7')
        self.assertEqual(page.status_code,200)
        self.assertIn('77 001',page.text)
        self.assertNotIn('998 877',page.text)
        with self.db:self.db.execute("UPDATE leads SET first_seen=(now()-interval '20 days')::text WHERE id=?",(self.ids[0],))
        self.assertNotIn('77 001',self.client.get('/app?days=7').text)
        self.assertIn('77 001',self.client.get('/app?days=30').text)
        self.assertEqual(self.client.get('/app?days=1000').status_code,400)

    def test_duplicate_does_not_send_again_to_user(self):
        from leadgen.product import deliver_registered
        from test_leadgen import CONFIG
        payload=json.loads(self.db.execute('SELECT payload FROM leads WHERE id=?',(self.ids[0],)).fetchone()['payload'])
        payload.update(url='https://t.me/another/1',title='Нужен бот',text='Нужен бот, бюджет 20000 рублей',budget_text='20000 руб.')
        with self.db:
            self.db.execute("UPDATE user_preferences SET monitoring_active=true,show_without_budget=true WHERE user_id=?",(self.u,))
            self.db.execute('INSERT INTO leads(url,fingerprint,payload,status,reason,first_seen) VALUES(?,?,?,?,?,?)',
                (payload['url'],'0',json.dumps(payload),'ready','test',datetime.now(timezone.utc).isoformat()))
        with self.db:self.db.execute("UPDATE leads SET status='duplicate' WHERE id=?",(self.ids[1],))
        sent=[]
        self.assertEqual(deliver_registered(self.db,CONFIG|{'sources':[]},lambda *a:sent.append(a)),0)
        self.assertEqual(sent,[])

    def test_personal_telegram_connection_and_group_selection(self):
        self.assertEqual(self.client.get('/app/telegram').status_code,200)
        response=self.client.post('/app/telegram/connect',data={'csrf':'test-csrf','phone':'+79991234567'})
        self.assertEqual(response.status_code,302)
        connection=self.db.execute('SELECT * FROM telegram_connections WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual(connection['phone_hint'],'+79••••••567')
        auth=self.db.execute('SELECT state_cipher FROM telegram_connection_auth WHERE user_id=?',(self.u,)).fetchone()
        self.assertNotIn('+79991234567',auth['state_cipher'])
        response=self.client.post('/app/telegram/code',data={'csrf':'test-csrf','code':'1111'})
        self.assertEqual(response.status_code,302)
        connection=self.db.execute('SELECT * FROM telegram_connections WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual(connection['status'],'active')
        self.assertNotIn('active-session',connection['session_cipher'])
        self.assertEqual(self.db.execute('SELECT count(*) AS n FROM user_telegram_dialogs WHERE user_id=?',(self.u,)).fetchone()['n'],2)
        self.client.post('/app/telegram/groups',data={'csrf':'test-csrf','groups':'-10001'})
        selected=self.db.execute('SELECT peer_id FROM user_telegram_dialogs WHERE user_id=? AND enabled=true',(self.u,)).fetchall()
        self.assertEqual([row['peer_id'] for row in selected],[-10001])
        self.assertEqual(self.client.post('/app/telegram/groups',data={'csrf':'test-csrf','groups':'-99999'}).status_code,302)
        self.assertFalse(self.db.execute('SELECT 1 FROM user_telegram_dialogs WHERE user_id=? AND enabled=true',(self.u,)).fetchone())
        response=self.client.post('/app/telegram/disconnect',data={'csrf':'test-csrf'})
        self.assertEqual(response.status_code,302)
        connection=self.db.execute('SELECT status,session_cipher FROM telegram_connections WHERE user_id=?',(self.u,)).fetchone()
        self.assertEqual((connection['status'],connection['session_cipher']),('revoked',''))

    def test_personal_telegram_two_factor_flow(self):
        self.client.post('/app/telegram/connect',data={'csrf':'test-csrf','phone':'+79991234567'})
        self.client.post('/app/telegram/code',data={'csrf':'test-csrf','code':'2222'})
        self.assertEqual(self.db.execute('SELECT stage FROM telegram_connection_auth WHERE user_id=?',(self.u,)).fetchone()['stage'],'password')
        self.client.post('/app/telegram/password',data={'csrf':'test-csrf','password':'wrong'})
        self.assertEqual(self.db.execute('SELECT status FROM telegram_connections WHERE user_id=?',(self.u,)).fetchone()['status'],'pending')
        self.client.post('/app/telegram/password',data={'csrf':'test-csrf','password':'correct'})
        self.assertEqual(self.db.execute('SELECT status FROM telegram_connections WHERE user_id=?',(self.u,)).fetchone()['status'],'active')

    def test_private_origin_is_delivered_only_to_connection_owner(self):
        from leadgen.product import deliver_registered
        from test_leadgen import CONFIG
        payload=dict(source='Telegram · Личная группа',url='https://t.me/c/1/2',title='Нужен сайт',
            text='Нужен сайт, бюджет 20000 рублей',published=datetime.now(timezone.utc).isoformat(),budget_text='20000 рублей')
        with self.db:
            self.db.execute('UPDATE user_preferences SET monitoring_active=true WHERE user_id IN (?,?)',(self.u,self.other))
            self.db.execute('INSERT INTO leads(url,fingerprint,payload,status,reason,first_seen,origin_user_id) VALUES(?,?,?,?,?,?,?)',
                (payload['url'],'private',json.dumps(payload),'ready','test',payload['published'],self.u))
        sent=[]
        deliver_registered(self.db,CONFIG|{'sources':[]},lambda method,data:sent.append(data),limit_per_user=10)
        self.assertEqual([item['chat_id'] for item in sent],['1'])

    def test_team_invite_roles_shared_leads_and_assignment(self):
        response=self.client.post('/app/team/invite',data={'csrf':'test-csrf','role':'manager'})
        self.assertEqual(response.status_code,302)
        with self.client.session_transaction() as state:
            invite_path=state['new_invite_path'];state['user_id']=self.other;state['workspace_id']=self.other_ws
        self.assertIn('Команда',self.client.get(invite_path).text)
        self.assertEqual(self.client.post(invite_path,data={'csrf':'test-csrf'}).status_code,302)
        membership=self.db.execute('SELECT role FROM workspace_members WHERE workspace_id=? AND user_id=?',
                                   (self.ws,self.other)).fetchone()
        self.assertEqual(membership['role'],'manager')
        self.assertEqual(self.client.get(f'/app/leads/{self.ids[0]}').status_code,200)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}',data={
            'csrf':'test-csrf','status':'working','amount':'','feedback':''}).status_code,302)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}/assign',data={
            'csrf':'test-csrf','assignee':str(self.other)}).status_code,302)
        assignment=self.db.execute('SELECT assignee_user_id FROM lead_assignments WHERE workspace_id=? AND lead_id=?',
                                   (self.ws,self.ids[0])).fetchone()
        self.assertEqual(assignment['assignee_user_id'],self.other)
        from leadgen.team_notifications import deliver_assignments
        sent=[]
        self.assertEqual(deliver_assignments(self.db,lambda method,data:sent.append(data)),1)
        self.assertEqual(sent[0]['chat_id'],'2')
        self.assertEqual(deliver_assignments(self.db,lambda method,data:sent.append(data)),0)
        with self.db:self.db.execute("UPDATE workspace_members SET role='viewer' WHERE workspace_id=? AND user_id=?",
                                    (self.ws,self.other))
        self.assertEqual(self.client.get(f'/app/leads/{self.ids[0]}').status_code,200)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}',data={
            'csrf':'test-csrf','status':'won','amount':'1','feedback':'fit'}).status_code,403)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}/notes',data={
            'csrf':'test-csrf','note':'Нельзя'}).status_code,403)
        self.assertEqual(self.client.post(invite_path,data={'csrf':'test-csrf'}).status_code,400)

    def test_admin_cannot_promote_admin_or_change_owner(self):
        with self.db:
            self.db.execute("INSERT INTO workspace_members(workspace_id,user_id,role) VALUES(?,?,'admin')",
                            (self.ws,self.other))
        with self.client.session_transaction() as state:
            state['user_id']=self.other;state['workspace_id']=self.ws
        self.assertEqual(self.client.post(f'/app/team/members/{self.u}',data={
            'csrf':'test-csrf','role':'viewer'}).status_code,403)
        self.assertEqual(self.client.post('/app/team/invite',data={
            'csrf':'test-csrf','role':'admin'}).status_code,403)

    def test_drafts_feedback_tags_and_bulk_actions(self):
        detail=self.client.get(f'/app/leads/{self.ids[0]}')
        for name in ('Короткий','Экспертный','Дружелюбный'):
            self.assertIn(name,detail.text)
        response=self.client.post('/app/leads/bulk',data={
            'csrf':'test-csrf','leads':str(self.ids[0]),'action':'feedback:competitor'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.db.execute('SELECT label FROM lead_feedback WHERE user_id=? AND lead_id=?',
                                        (self.u,self.ids[0])).fetchone()['label'],'competitor')
        response=self.client.post('/app/leads/bulk',data={
            'csrf':'test-csrf','leads':str(self.ids[0]),'action':'tag','tag':'важный'})
        self.assertEqual(response.status_code,302)
        self.assertIn('важный',self.client.get(f'/app/leads/{self.ids[0]}').text)
        self.assertEqual(self.client.post('/app/leads/bulk',data={
            'csrf':'test-csrf','leads':str(self.ids[1]),'action':'status:won'}).status_code,404)
        self.assertEqual(self.client.post(f'/app/leads/{self.ids[0]}/tags',data={
            'csrf':'test-csrf','tag':'важный','action':'remove'}).status_code,302)
        self.assertFalse(self.db.execute('SELECT 1 FROM lead_tags WHERE user_id=? AND lead_id=?',
                                         (self.u,self.ids[0])).fetchone())

    def test_api_tokens_and_signed_webhook_delivery(self):
        import hashlib,hmac
        from leadgen.integrations import deliver_webhooks
        response=self.client.post('/app/integrations/tokens',data={
            'csrf':'test-csrf','name':'CRM чтение','access':'read'})
        self.assertEqual(response.status_code,302)
        with self.client.session_transaction() as state:read_token=state['new_api_token']
        headers={'Authorization':'Bearer '+read_token}
        payload=self.client.get('/api/v1/leads',headers=headers).get_json()
        self.assertEqual([row['id'] for row in payload['data']],[self.ids[0]])
        self.assertEqual(self.client.patch(f'/api/v1/leads/{self.ids[0]}',headers=headers,
            json={'status':'won'}).status_code,403)
        self.client.post('/app/integrations/tokens',data={'csrf':'test-csrf','name':'CRM запись','access':'write'})
        with self.client.session_transaction() as state:write_token=state['new_api_token']
        response=self.client.patch(f'/api/v1/leads/{self.ids[0]}',headers={
            'Authorization':'Bearer '+write_token},json={'status':'won','deal_amount':12345,'feedback':'fit'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.db.execute('SELECT deal_amount FROM user_leads WHERE user_id=? AND lead_id=?',
                                        (self.u,self.ids[0])).fetchone()['deal_amount'],12345)
        self.client.post('/app/integrations/webhooks',data={'csrf':'test-csrf','name':'CRM','url':'https://example.com/hook'})
        with self.client.session_transaction() as state:secret=state['new_webhook_secret']
        stored=self.db.execute('SELECT secret_cipher FROM webhooks WHERE workspace_id=?',(self.ws,)).fetchone()['secret_cipher']
        self.assertNotIn(secret,stored)
        self.client.post(f'/app/leads/{self.ids[0]}/tags',data={
            'csrf':'test-csrf','tag':'передать-в-crm','action':'add'})
        queued=self.db.execute('SELECT payload FROM webhook_outbox ORDER BY id DESC LIMIT 1').fetchone()['payload']
        self.assertIn('передать-в-crm',queued['lead']['tags'])
        self.assertEqual(queued['lead']['last_activity'],None)
        calls=[]
        def sender(url,body,request_headers,resolver):calls.append((url,body,request_headers))
        self.assertEqual(deliver_webhooks(self.db,self.app.config['TELEGRAM_CIPHER_KEY'],sender,
            self.app.config['WEBHOOK_RESOLVER']),1)
        self.assertEqual(calls[0][0],'https://example.com/hook')
        signature='sha256='+hmac.new(secret.encode(),calls[0][1],hashlib.sha256).hexdigest()
        self.assertEqual(calls[0][2]['X-Leadfinder-Signature'],signature)
        self.assertEqual(deliver_webhooks(self.db,self.app.config['TELEGRAM_CIPHER_KEY'],sender,
            self.app.config['WEBHOOK_RESOLVER']),0)
        self.assertEqual(self.client.get('/api/v1/leads',headers={'Authorization':'Bearer wrong'}).status_code,401)

    def test_demand_report_is_scoped_and_period_bound(self):
        page=self.client.get('/app/demand?days=30')
        self.assertEqual(page.status_code,200)
        self.assertIn('Всего запросов</span><strong>1',page.text)
        self.assertIn('Сайты',page.text)
        with self.db:self.db.execute("UPDATE leads SET first_seen=(now()-interval '20 days')::text WHERE id=?",(self.ids[0],))
        self.assertIn('Всего запросов</span><strong>0',self.client.get('/app/demand?days=7').text)
        self.assertIn('Всего запросов</span><strong>1',self.client.get('/app/demand?days=30').text)
        self.assertEqual(self.client.get('/app/demand?days=365').status_code,400)

    def test_ai_offers_are_on_demand_persisted_scoped_and_rate_limited(self):
        with patch.dict(os.environ,{'GROQ_API_KEY':'test-key','GROQ_MODEL':'openai/gpt-oss-120b'}):
            response=self.client.post(f'/app/leads/{self.ids[0]}/ai-offers',data={'csrf':'test-csrf','project_id':''})
            self.assertEqual(response.status_code,302)
            page=self.client.get(f'/app/leads/{self.ids[0]}')
            self.assertIn('Готов обсудить вашу задачу',page.text)
            record=self.db.execute('SELECT * FROM ai_offer_generations WHERE owner_user_id=?',(self.u,)).fetchone()
            self.assertEqual(record['status'],'completed')
            self.assertEqual((record['input_tokens'],record['output_tokens']),(321,123))
            self.client.post(f'/app/leads/{self.ids[0]}/ai-offers',data={'csrf':'test-csrf','project_id':''})
            self.assertEqual(self.db.execute('SELECT count(*) AS n FROM ai_offer_generations WHERE owner_user_id=?',
                                             (self.u,)).fetchone()['n'],1)
            self.assertEqual(self.client.post(f'/app/leads/{self.ids[1]}/ai-offers',data={
                'csrf':'test-csrf','project_id':''}).status_code,404)
