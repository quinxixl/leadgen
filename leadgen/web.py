"""Server-rendered customer portal. All customer queries are owner scoped."""
import csv
import io
import json
import os
import re
import secrets
import hmac
from datetime import date, datetime, timedelta, timezone
from urllib.parse import quote, urlsplit

from flask import Flask, abort, current_app, flash, g, has_request_context, jsonify, redirect, render_template, request, send_from_directory, session, url_for, Response
from flask.sessions import SecureCookieSessionInterface
from itsdangerous import BadSignature
from werkzeug.middleware.proxy_fix import ProxyFix
from .app import env_load
from .database import db_open
from .core import TOPICS, TOPIC_CATEGORIES, Lead, enrich
from .product import PIPELINE, FEEDBACK, ProductController
from .web_auth import begin_login, consume_handoff, consume_login, login_status
from .billing import PLANS, PLAN_NAMES, days_left, legal_entity, subscription_active, support_contact
from .rate_limit import RateLimited, hit, purge_expired


def metrika_id():
    """Yandex Metrica counter number; empty disables the counter."""
    value=os.environ.get('YANDEX_METRIKA_ID','').strip()
    return value if value.isdigit() else ''


def yandex_verification():
    """Site-verification token for Яндекс.Вебмастер; empty disables the meta tag."""
    return os.environ.get('YANDEX_VERIFICATION','').strip()


def google_verification():
    """Site-verification token for Google Search Console; empty disables the meta tag."""
    return os.environ.get('GOOGLE_VERIFICATION','').strip()


def platform_admin_ids():
    return {value.strip() for value in os.environ.get('WEB_ADMIN_TELEGRAM_IDS','').split(',') if value.strip()}


def is_platform_admin(user):
    return bool(user and str(user['telegram_user_id']) in platform_admin_ids())


# Paths that stay open when the subscription has ended, so the user can pay or leave.
BILLING_OPEN=('/app/billing','/app/workspace','/app/telegram/disconnect','/app/team/join/','/app/security',
              '/tg/billing','/tg/handoff')
# Mini App paths reachable without a session: the boot page and the initData exchange.
TG_PUBLIC=('/tg','/tg/auth')
TG_SESSION_LIFETIME=timedelta(hours=4)
SAFE_PATH=re.compile(r'/[A-Za-z0-9_\-./]*(?:\?[A-Za-z0-9_\-.=&%]*)?')


def is_tg_path(path):
    return path=='/tg' or path.startswith('/tg/')


def safe_next(value,prefix):
    """A same-site path under prefix ('/app' or '/tg'), or '' for anything that could leave the site."""
    value=value if isinstance(value,str) else ''
    if (len(value)>300 or not SAFE_PATH.fullmatch(value) or '//' in value or '..' in value
            or not value.startswith(prefix) or (len(value)>len(prefix) and value[len(prefix)] not in '/?')):
        return ''
    return value


class _TgSigner(SecureCookieSessionInterface):
    salt='signalid-tg-session'


class SplitSessionInterface(SecureCookieSessionInterface):
    """The Mini App (/tg) keeps its own cookie. Telegram Web embeds it in a cross-site iframe, so the
    cookie must be SameSite=None and Partitioned (CHIPS); the cabinet cookie stays SameSite=Lax."""
    tg_signer=_TgSigner()

    @staticmethod
    def _tg():
        return has_request_context() and is_tg_path(request.path)

    def get_signing_serializer(self,app):
        return self.tg_signer.get_signing_serializer(app) if self._tg() else super().get_signing_serializer(app)

    def get_cookie_name(self,app):
        return 'signalid_tg' if self._tg() else super().get_cookie_name(app)

    def get_cookie_path(self,app):
        return '/tg' if self._tg() else super().get_cookie_path(app)

    def get_cookie_samesite(self,app):
        return 'None' if self._tg() else super().get_cookie_samesite(app)

    def get_cookie_secure(self,app):
        return True if self._tg() else super().get_cookie_secure(app)

    def get_cookie_partitioned(self,app):
        return True if self._tg() else super().get_cookie_partitioned(app)

    def get_expiration_time(self,app,session):
        if self._tg() and session.permanent:
            return datetime.now(timezone.utc)+TG_SESSION_LIFETIME
        return super().get_expiration_time(app,session)

    def open_session(self,app,request):
        if not is_tg_path(request.path):
            return super().open_session(app,request)
        signer=self.get_signing_serializer(app)
        if signer is None:return None
        value=request.cookies.get('signalid_tg')
        if not value:return self.session_class()
        try:return self.session_class(signer.loads(value,max_age=int(TG_SESSION_LIFETIME.total_seconds())))
        except BadSignature:return self.session_class()


def owner_lead(db,owner_id,lid):
    """The owner's delivered lead with CRM state, or None."""
    return db.execute('''SELECT l.*,ul.pipeline_status,ul.deal_amount,ul.filter_reason,
        f.label AS feedback FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
        LEFT JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
        WHERE ul.user_id=? AND l.id=? AND ul.delivery_status<>'filtered' ''',(owner_id,lid)).fetchone()


def update_lead(db,owner_id,lid,row,status,amount,label):
    """Save status, deal amount (digits or '') and feedback with activity rows and a webhook event.

    Shared by the cabinet and the Mini App; callers validate the values first."""
    with db:
        for kind,value,old in [('status',status,row['pipeline_status']),('amount',amount,str(row['deal_amount']) if row['deal_amount'] is not None else ''),('feedback',label,row['feedback'] or '')]:
            if value!=old:
                detail=PIPELINE.get(value,value) if kind=='status' else FEEDBACK.get(value,value) if kind=='feedback' else value or 'Сумма не указана'
                db.execute('INSERT INTO lead_activity(user_id,lead_id,kind,detail) VALUES(?,?,?,?)',(owner_id,lid,kind,detail or 'Без оценки'))
        db.execute('UPDATE user_leads SET pipeline_status=?,deal_amount=?,updated_at=now() WHERE user_id=? AND lead_id=?',
            (status,int(amount) if amount else None,owner_id,lid))
        if label:db.execute('''INSERT INTO lead_feedback(user_id,lead_id,label) VALUES(?,?,?)
            ON CONFLICT(user_id,lead_id) DO UPDATE SET label=excluded.label,updated_at=now()''',(owner_id,lid,label))
        else:db.execute('DELETE FROM lead_feedback WHERE user_id=? AND lead_id=?',(owner_id,lid))
        from .integrations import enqueue_lead_event
        enqueue_lead_event(db,owner_id,lid,'lead.updated')


PREFERENCE_COLUMNS={'min_budget':'?','topics':'?::jsonb','show_without_budget':'?','show_possible_needs':'?',
                    'monitoring_active':'?','portfolio':'?'}


def save_preferences(db,owner_id,**values):
    """Update the owner's lead filters; leads filtered out earlier are re-evaluated under the new rules."""
    columns=','.join(f'{name}={PREFERENCE_COLUMNS[name]}' for name in values)
    params=[json.dumps(value) if name=='topics' else value for name,value in values.items()]
    with db:
        db.execute(f'UPDATE user_preferences SET {columns},updated_at=now() WHERE user_id=?',params+[owner_id])
        db.execute("DELETE FROM user_leads WHERE user_id=? AND delivery_status='filtered'",(owner_id,))


def start_user_session(db,user_id,permanent=True):
    """Replace the current cookie session with a signed-in one bound to the user's session_version."""
    row=db.execute('SELECT session_version FROM app_users WHERE id=?',(user_id,)).fetchone()
    session.clear()
    session.update(user_id=user_id,sv=row['session_version'] if row else 0,csrf=secrets.token_urlsafe(32))
    session.permanent=permanent


def like_pattern(value):
    """Literal substring for ILIKE: user input must not act as a wildcard."""
    return '%'+value.replace('\\','\\\\').replace('%','\\%').replace('_','\\_')+'%'


def browser_hint(user_agent, now=None):
    """Short, non-identifying description of the browser asking to sign in."""
    ua=user_agent or ''
    browser=next((name for key,name in (('YaBrowser','Яндекс Браузер'),('Edg/','Edge'),('OPR/','Opera'),
        ('Firefox/','Firefox'),('Chrome/','Chrome'),('Safari/','Safari')) if key in ua),'браузер')
    system=next((name for key,name in (('Android','Android'),('iPhone','iPhone'),('iPad','iPad'),
        ('Windows','Windows'),('Mac OS X','macOS'),('Linux','Linux')) if key in ua),'')
    from datetime import datetime,timezone
    moment=(now or datetime.now(timezone.utc)).astimezone(timezone(timedelta(hours=3))).strftime('%H:%M МСК')
    return ' · '.join(x for x in (browser,system,moment) if x)


def create_app(test_config=None):
    env_load()
    app = Flask(__name__, template_folder='web_templates', static_folder='web_static')
    # compose.web.yaml exposes this process only on loopback behind one HTTPS proxy.
    app.wsgi_app=ProxyFix(app.wsgi_app,x_for=1,x_proto=1,x_host=1)
    app.session_interface=SplitSessionInterface()
    app.config.update(SECRET_KEY=os.environ.get('WEB_SECRET_KEY'),
        SESSION_COOKIE_NAME='signalid_session', SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SECURE=True, SESSION_COOKIE_SAMESITE='Lax',
        # Long-lived cabinet sessions; «Выйти на всех устройствах» revokes them via session_version.
        PERMANENT_SESSION_LIFETIME=timedelta(days=30), MAX_CONTENT_LENGTH=65536,
        BOT_USERNAME=os.environ.get('TELEGRAM_BOT_USERNAME',''),
        TELEGRAM_BOT_TOKEN=os.environ.get('TELEGRAM_BOT_TOKEN',''),
        WEB_PUBLIC_URL=os.environ.get('WEB_PUBLIC_URL','').strip().rstrip('/'),
        TELEGRAM_CIPHER_KEY=os.environ.get('TELEGRAM_SESSION_ENCRYPTION_KEY',''),
        DB_FACTORY=db_open, TRUSTED_HOSTS=os.environ.get('WEB_ALLOWED_HOSTS','localhost,127.0.0.1').split(','))
    if test_config:app.config.update(test_config)
    if not app.config['SECRET_KEY'] or len(app.config['SECRET_KEY']) < 32:
        raise ValueError('WEB_SECRET_KEY должен содержать минимум 32 символа')

    def db():
        if 'db' not in g:g.db=app.config['DB_FACTORY']()
        return g.db

    @app.teardown_appcontext
    def close(error=None):
        connection=g.pop('db',None)
        if connection:connection.close()

    @app.before_request
    def protect():
        tg=is_tg_path(request.path)
        session.setdefault('csrf',secrets.token_urlsafe(32))
        if request.method=='POST' and not request.path.startswith('/api/'):
            if request.path=='/tg/auth':
                # Telegram signs initData, which replaces the form token; cross-site posts are still refused.
                origin=request.headers.get('Origin')
                if origin and origin!=request.host_url.rstrip('/'):abort(403)
            elif not hmac.compare_digest(session['csrf'].encode(),request.form.get('csrf','').encode()):
                abort(400,description='Сессия формы истекла. Обновите страницу.')
        g.user=None
        if session.get('user_id'):
            user=db().execute('SELECT * FROM app_users WHERE id=? AND status=\'active\'',(session['user_id'],)).fetchone()
            # «Выйти на всех устройствах» bumps session_version; older cookies stop working at once.
            if user and session.get('sv',0)==user['session_version']:g.user=user
            else:
                session.clear();session['csrf']=secrets.token_urlsafe(32)
        if request.path.startswith(('/app','/admin')) and not g.user:
            return redirect(url_for('login'))
        g.workspaces=[];g.workspace=None;g.role=None;g.owner_id=None
        tg_private=tg and request.path not in TG_PUBLIC
        if tg_private and not g.user:
            # Opened as a Mini App URL without a cookie: the boot page signs in with initData and comes back.
            # The browser keeps the #tgWebAppData fragment across this redirect.
            if request.method=='GET':
                return redirect('/tg?next='+quote(request.full_path.rstrip('?'),safe='/'))
            return redirect('/tg')
        if g.user and (request.path.startswith('/app') or tg_private):
            from .teams import load_workspace
            g.workspaces,g.workspace=load_workspace(db(),g.user,session.get('workspace_id'))
            session['workspace_id']=g.workspace['id'];g.role='owner' if is_platform_admin(g.user) else g.workspace['role'];g.owner_id=g.workspace['owner_user_id']
            # Team members work under the workspace owner's subscription.
            g.subscription=db().execute('SELECT status,plan_code,ends_at FROM subscriptions WHERE user_id=?',(g.owner_id,)).fetchone()
            if tg and not g.user['terms_accepted_at']:
                # Consent is the first onboarding step; nothing else in the Mini App opens before it.
                if request.path!='/tg/start':
                    target=request.full_path.rstrip('?') if request.method=='GET' and request.path!='/tg/feed' else ''
                    return redirect('/tg/start'+('?next='+quote(target,safe='/') if target else ''))
                return
            if (not subscription_active(g.subscription) and not is_platform_admin(g.user)
                    and not request.path.startswith(BILLING_OPEN)):
                return redirect('/tg/billing' if tg else url_for('billing'))

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Referrer-Policy']='same-origin'
        if request.is_secure:
            response.headers['Strict-Transport-Security']='max-age=31536000; includeSubDomains'
        response.headers['Permissions-Policy']='camera=(), microphone=(), geolocation=(), payment=()'
        if is_tg_path(request.path):
            # Mini App: Telegram's SDK script, embedding by Telegram Web only, no analytics. The SDK injects
            # Telegram Web's iframe styles into a <style> element, hence inline styles here and only here.
            response.headers['Content-Security-Policy']=("default-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; "
                "script-src 'self' https://telegram.org; connect-src 'self'; frame-ancestors https://web.telegram.org; "
                "base-uri 'self'; form-action 'self'")
            response.headers['Cache-Control']='no-store'
            return response
        # Yandex Metrica runs on public pages only: the cabinet holds client data that must not reach analytics.
        public=metrika_id() and not request.path.startswith(('/app','/admin','/api/','/login'))
        yandex=' https://mc.yandex.ru https://mc.yandex.com https://yastatic.net' if public else ''
        # Metrica's click map shows public pages in a frame on metrika.yandex.ru; everything else stays unframeable.
        ancestors="'self' https://metrika.yandex.ru https://metrika.yandex.by https://metrica.yandex.com https://*.webvisor.com" if public else "'none'"
        if not public:response.headers['X-Frame-Options']='DENY'
        # Metrica's element picker and click map draw overlays with inline styles and blob frames.
        styles="'self' 'unsafe-inline'" if public else "'self'"
        frames=f"; frame-src blob:{yandex}" if public else ''
        response.headers['Content-Security-Policy']=(f"default-src 'self'; style-src {styles}; img-src 'self' data:{yandex}; "
            f"script-src 'self'{yandex}; connect-src 'self'{yandex}{frames}; frame-ancestors {ancestors}; base-uri 'self'; form-action 'self'")
        if not request.path.startswith('/web_static/'):
            response.headers['Cache-Control']='no-store'
        return response

    @app.context_processor
    def context():
        from .teams import ROLES
        sub=g.get('subscription')
        return dict(csrf=session.get('csrf',''),user=g.get('user'),pipeline=PIPELINE,feedback=FEEDBACK,
                    subscription=sub,subscription_ok=subscription_active(sub),plan_names=PLAN_NAMES,
                    trial_days_left=days_left(sub) if sub else None,support=support_contact(),plans=PLANS,
                    legal_entity=legal_entity(),metrika_id=metrika_id(),
                    yandex_verification=yandex_verification(),google_verification=google_verification(),
                    workspaces=g.get('workspaces',[]),workspace=g.get('workspace'),role=g.get('role'),roles=ROLES,
                    platform_admin=is_platform_admin(g.get('user')))

    @app.template_filter('original_url')
    def original_url(value):
        parts=urlsplit(value or '')
        return value if parts.scheme=='https' and parts.hostname else '#'

    @app.template_filter('dt')
    def moscow_time(value):
        """ISO string or datetime as «03.10.2026 18:19 МСК»; unknown values pass through."""
        from datetime import datetime,timezone
        if not value:return ''
        try:moment=value if isinstance(value,datetime) else datetime.fromisoformat(str(value))
        except ValueError:return value
        if moment.tzinfo is None:moment=moment.replace(tzinfo=timezone.utc)
        return moment.astimezone(timezone(timedelta(hours=3))).strftime('%d.%m.%Y %H:%M МСК')

    @app.template_filter('money')
    def money(value):return f'{value:,}'.replace(',',' ') if value is not None else 'Не указана'

    @app.get('/')
    def landing():return render_template('landing.html')

    @app.get('/pricing')
    def pricing():return render_template('pricing.html')

    @app.get('/privacy')
    def privacy():return render_template('privacy.html')

    @app.get('/terms')
    def terms():return render_template('terms.html')

    @app.get('/favicon.ico')
    def favicon():
        # Search engines (notably Yandex) fetch /favicon.ico from the site root.
        return send_from_directory(app.static_folder, 'favicon.ico', mimetype='image/x-icon')

    @app.get('/robots.txt')
    def robots():
        sitemap_url=request.url_root.rstrip('/')+'/sitemap.xml'
        return Response('User-agent: *\nDisallow: /app\nDisallow: /admin\nDisallow: /login\n'
            f'Sitemap: {sitemap_url}\n',mimetype='text/plain')

    @app.get('/sitemap.xml')
    def sitemap():
        root=request.url_root.rstrip('/')
        # Indexable public pages only; /login and /app are disallowed in robots.txt.
        locs=''.join(f'<url><loc>{root}{path}</loc></url>' for path in ('/','/pricing','/privacy','/terms'))
        xml=('<?xml version="1.0" encoding="UTF-8"?>\n'
            f'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">{locs}</urlset>')
        return Response(xml,mimetype='application/xml')

    @app.get('/app/billing')
    def billing():
        return render_template('billing.html',active=subscription_active(g.subscription))

    @app.get('/healthz')
    def health():return {'status':'ok'}

    @app.route('/login',methods=['GET','POST'])
    def login():
        if g.user:return redirect(url_for('dashboard'))
        if not app.config['BOT_USERNAME']:
            return render_template('error.html',message='Вход ещё не настроен администратором.'),503
        if request.method=='POST':
            if request.form.get('action')=='finish':
                try:
                    uid=consume_login(db(),session.get('login_token',''),session.get('login_browser',''))
                    if uid:
                        # A login request can only be started with the consent checkbox ticked (checked above).
                        from .accounts import accept_terms
                        accept_terms(db(),uid)
                        start_user_session(db(),uid)
                        return redirect(url_for('dashboard'))
                    flash('Сначала подтвердите вход в Telegram.')
                except ValueError as exc:
                    flash(str(exc));session.pop('login_token',None)
            else:
                if not request.form.get('consent'):
                    flash('Чтобы продолжить, примите условия оферты и согласие на обработку персональных данных.')
                    return render_template('login.html',bot=app.config['BOT_USERNAME']),400
                try:
                    hit(db(),'login-ip',request.remote_addr or '',10,600)
                except RateLimited as exc:
                    flash(str(exc));return render_template('login.html',bot=app.config['BOT_USERNAME']),429
                if secrets.randbelow(20)==0:purge_expired(db())
                token,browser,code=begin_login(db(),browser_hint(request.headers.get('User-Agent')))
                session.update(login_token=token,login_browser=browser,login_code=code)
        return render_template('login.html',bot=app.config['BOT_USERNAME'])

    @app.get('/login/status')
    def login_poll():
        """Lets the login page finish on its own once the bot approves the code."""
        token=session.get('login_token','')
        if g.user:return jsonify(status='signed_in')
        if not token:return jsonify(status='none')
        return jsonify(status=login_status(db(),token,session.get('login_browser','')))

    @app.get('/login/handoff')
    def login_handoff():
        """One-time link from the Mini App: opens the cabinet in the external browser already signed in."""
        target=safe_next(request.args.get('next','/app'),'/app') or '/app'
        try:
            uid,version=consume_handoff(db(),request.args.get('t',''))
        except ValueError as exc:
            flash(str(exc));return redirect(url_for('login'))
        user=db().execute("SELECT id,session_version FROM app_users WHERE id=? AND status='active'",(uid,)).fetchone()
        if not user or user['session_version']!=version:
            flash('Ссылка входа устарела. Откройте сайт из Telegram ещё раз.');return redirect(url_for('login'))
        start_user_session(db(),uid)
        return redirect(target)

    @app.post('/logout')
    def logout():session.clear();return redirect(url_for('landing'))

    @app.get('/app/security')
    def security():
        return render_template('security.html')

    @app.post('/app/security/logout-all')
    def logout_everywhere():
        with db():db().execute('UPDATE app_users SET session_version=session_version+1 WHERE id=?',(g.user['id'],))
        session.clear()
        flash('Вы вышли на всех устройствах. Войдите снова.')
        return redirect(url_for('login'))

    def user_lead(lid):
        row=owner_lead(db(),g.owner_id,lid)
        if not row:abort(404)
        return row

    @app.get('/app')
    def dashboard():
        try:days=int(request.args.get('days','30'))
        except ValueError:abort(400)
        if days not in (7,30,90):abort(400)
        owner=(g.owner_id,days)
        base="ul.user_id=? AND ul.delivery_status<>'filtered' AND ul.updated_at IS NOT NULL AND l.first_seen::timestamptz >= now()-(? * interval '1 day')"
        stats=db().execute("""SELECT count(*) AS total,
            count(*) FILTER (WHERE ul.pipeline_status='won') AS won,
            count(*) FILTER (WHERE f.label='fit') AS fit,
            sum(ul.deal_amount) FILTER (WHERE ul.pipeline_status='won') AS revenue
            FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
            LEFT JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
            WHERE """+base,owner).fetchone()
        stages=db().execute("SELECT ul.pipeline_status,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE "+base+" GROUP BY ul.pipeline_status",owner).fetchall()
        sources=db().execute("""SELECT l.payload::jsonb->>'source' AS name,count(*) AS count,
            count(f.label) AS reviewed,count(*) FILTER (WHERE f.label='fit') AS fit,
            count(*) FILTER (WHERE f.label='ad') AS ads,
            count(*) FILTER (WHERE ul.pipeline_status='won') AS won
            FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
            LEFT JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
            WHERE """+base+" GROUP BY name ORDER BY won DESC,fit DESC,count DESC LIMIT 20",owner).fetchall()
        daily=db().execute("SELECT (l.first_seen::timestamptz AT TIME ZONE 'UTC')::date AS day,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE "+base+" GROUP BY day ORDER BY day",owner).fetchall()
        quality=db().execute("""SELECT f.label,count(*) AS count FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
            JOIN lead_feedback f ON f.user_id=ul.user_id AND f.lead_id=ul.lead_id
            WHERE """+base+" GROUP BY f.label ORDER BY count DESC",owner).fetchall()
        response=db().execute("""SELECT avg(extract(epoch FROM (a.first_response-ul.sent_at)))/60 AS minutes,count(a.first_response) AS samples
            FROM user_leads ul JOIN leads l ON l.id=ul.lead_id
            JOIN LATERAL (SELECT min(created_at) AS first_response FROM lead_activity
                WHERE user_id=ul.user_id AND lead_id=ul.lead_id AND kind='status' AND detail='Написали') a ON true
            WHERE """+base+" AND ul.sent_at IS NOT NULL AND a.first_response>=ul.sent_at",owner).fetchone()
        projects=db().execute("""SELECT p.name,count(*) AS count,count(*) FILTER (WHERE ul.pipeline_status='won') AS won
            FROM project_leads pl JOIN projects p ON p.id=pl.project_id AND p.user_id=pl.user_id
            JOIN user_leads ul ON ul.user_id=pl.user_id AND ul.lead_id=pl.lead_id JOIN leads l ON l.id=ul.lead_id
            WHERE """+base+" GROUP BY p.id ORDER BY count DESC",owner).fetchall()
        team=db().execute('''SELECT u.display_name,m.role,count(a.lead_id) AS assigned,
            count(a.lead_id) FILTER(WHERE ul.pipeline_status='won') AS won
            FROM workspace_members m JOIN app_users u ON u.id=m.user_id
            LEFT JOIN lead_assignments a ON a.workspace_id=m.workspace_id AND a.assignee_user_id=m.user_id
            LEFT JOIN user_leads ul ON ul.user_id=a.owner_user_id AND ul.lead_id=a.lead_id
            WHERE m.workspace_id=? GROUP BY u.id,m.role ORDER BY won DESC,assigned DESC,u.display_name''',
            (g.workspace['id'],)).fetchall()
        return render_template('dashboard.html',stats=stats,stages=stages,sources=sources,days=days,daily=daily,
            peak=max((r['count'] for r in daily),default=1),quality=quality,response=response,projects=projects,team=team)

    def filtered_leads(export=False):
        try:page=max(1,min(10000,int(request.args.get('page','1'))))
        except ValueError:abort(400)
        search=request.args.get('q','')[:200];status=request.args.get('status','')
        project=request.args.get('project','')
        args=[g.owner_id];where="ul.user_id=? AND ul.delivery_status<>'filtered'"
        if project:
            if not project.isdigit():abort(400)
            where+=' AND EXISTS (SELECT 1 FROM project_leads pl WHERE pl.user_id=ul.user_id AND pl.lead_id=l.id AND pl.project_id=?)'
            args.append(int(project))
        source=request.args.get('source','')[:300]
        label=request.args.get('feedback','')
        if source:
            where+=" AND l.payload::jsonb->>'source'=?";args.append(source)
        if label:
            if label not in FEEDBACK:abort(400)
            where+=' AND EXISTS (SELECT 1 FROM lead_feedback f WHERE f.user_id=ul.user_id AND f.lead_id=l.id AND f.label=?)';args.append(label)
        dates={}
        for key,operator in [('from','>='),('to','<')]:
            value=request.args.get(key,'')
            if value:
                try:day=date.fromisoformat(value)
                except ValueError:abort(400)
                dates[key]=day
                if key=='to':
                    if day==date.max:abort(400)
                    day+=timedelta(days=1)
                where+=' AND l.first_seen::timestamptz '+operator+' ?::timestamptz'
                args.append(day.isoformat()+'T00:00:00+00:00')
        if len(dates)==2 and dates['from']>dates['to']:abort(400)
        if search:where+=" AND l.payload ILIKE ? ESCAPE '\\'";args.append(like_pattern(search))
        if status:
            if status not in PIPELINE:abort(400)
            where+=' AND ul.pipeline_status=?';args.append(status)
        count=db().execute('SELECT count(*) AS n FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE '+where,args).fetchone()['n']
        if export and count>10000:abort(400,description='Сузьте выборку до 10 000 лидов для экспорта.')
        rows=db().execute('''SELECT l.id,l.payload,ul.pipeline_status,ul.deal_amount
            FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE '''+where+
            ' ORDER BY l.id DESC LIMIT ? OFFSET ?',args+([10000,0] if export else [30,(page-1)*30])).fetchall()
        return [dict(row)|{'lead':enrich(Lead(**json.loads(row['payload'])))} for row in rows],count,page

    @app.get('/app/leads')
    def leads():
        rows,count,page=filtered_leads()
        projects=db().execute('SELECT id,name FROM projects WHERE workspace_id=? ORDER BY name',(g.workspace['id'],)).fetchall()
        sources=db().execute("SELECT DISTINCT l.payload::jsonb->>'source' AS name FROM user_leads ul JOIN leads l ON l.id=ul.lead_id WHERE ul.user_id=? AND ul.delivery_status<>'filtered' ORDER BY name",(g.owner_id,)).fetchall()
        members=db().execute('''SELECT m.user_id,u.display_name FROM workspace_members m JOIN app_users u ON u.id=m.user_id
            WHERE m.workspace_id=? ORDER BY u.display_name''',(g.workspace['id'],)).fetchall()
        filters={k:request.args.get(k,'') for k in ('q','status','project','source','feedback','from','to')}
        return render_template('leads.html',rows=rows,count=count,page=page,projects=projects,sources=sources,
                               members=members,filters=filters)

    @app.get('/app/leads/export')
    def export():
        rows,_,_=filtered_leads(True)
        headings=['Название','Источник','Ссылка','Дата','Оценка','Статус','Сумма сделки']
        values=[]
        def safe(v):
            v=str(v or '')
            return "'"+v if v.lstrip().startswith(('=','+','-','@','\t','\r')) else v
        for r in rows:
            l=r['lead'];values.append([safe(x) for x in [l.title,l.source,l.url,l.published,l.score,PIPELINE[r['pipeline_status']],r['deal_amount']]])
        if request.args.get('format')=='xlsx':
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
            from openpyxl.utils import get_column_letter
            book=Workbook();sheet=book.active;sheet.title='Лиды';sheet.append(headings)
            for values_row in values:
                sheet.append([ILLEGAL_CHARACTERS_RE.sub('',v)[:32767] for v in values_row])
            for cell in sheet[1]:
                cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='08182F')
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for i,width in enumerate((55,32,45,28,12,24,20),1):sheet.column_dimensions[get_column_letter(i)].width=width
            for cells in sheet.iter_rows(min_row=2):
                for cell in cells:cell.alignment=Alignment(vertical='top',wrap_text=True)
            buffer=io.BytesIO();book.save(buffer)
            return Response(buffer.getvalue(),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename=leads.xlsx'})
        stream=io.StringIO();writer=csv.writer(stream,delimiter=';');writer.writerow(headings);writer.writerows(values)
        return Response('\ufeff'+stream.getvalue(),mimetype='text/csv',headers={'Content-Disposition':'attachment; filename=leads.csv'})

    @app.route('/app/leads/<int:lid>',methods=['GET','POST'])
    def detail(lid):
        from .teams import EDIT_LEADS,require_role
        row=user_lead(lid)
        if request.method=='POST':
            require_role(*EDIT_LEADS)
            status=request.form.get('status');label=request.form.get('feedback','')
            if status not in PIPELINE or label and label not in FEEDBACK:abort(400)
            amount=request.form.get('amount','').strip()
            if amount and (not amount.isdigit() or int(amount)>1000000000):
                abort(400,description='Укажите целую сумму от 0 до 1 000 000 000 ₽.')
            update_lead(db(),g.owner_id,lid,row,status,amount,label)
            flash('Изменения сохранены. Статус доступен и в Telegram.');return redirect(url_for('detail',lid=lid))
        activity=db().execute('SELECT kind,detail,created_at FROM lead_activity WHERE user_id=? AND lead_id=? ORDER BY id DESC LIMIT 100',(g.owner_id,lid)).fetchall()
        reminder=db().execute('SELECT * FROM lead_reminders WHERE user_id=? AND lead_id=?',(g.owner_id,lid)).fetchone()
        members=db().execute('''SELECT m.user_id,u.display_name FROM workspace_members m JOIN app_users u ON u.id=m.user_id
            WHERE m.workspace_id=? ORDER BY u.display_name''',(g.workspace['id'],)).fetchall()
        assignment=db().execute('''SELECT a.assignee_user_id,u.display_name FROM lead_assignments a
            JOIN app_users u ON u.id=a.assignee_user_id WHERE a.workspace_id=? AND a.lead_id=?''',
            (g.workspace['id'],lid)).fetchone()
        draft_projects=db().execute('''SELECT p.* FROM project_leads pl JOIN projects p
            ON p.id=pl.project_id AND p.user_id=pl.user_id
            WHERE pl.lead_id=? AND p.workspace_id=? ORDER BY p.id''',(lid,g.workspace['id'])).fetchall()
        requested_project=request.args.get('reply_project','')
        if requested_project and not requested_project.isdigit():abort(400)
        draft_project=next((item for item in draft_projects if str(item['id'])==requested_project),
                           draft_projects[0] if draft_projects else None)
        if requested_project and not draft_project:abort(404)
        profile=db().execute('SELECT profile_services,portfolio FROM user_preferences WHERE user_id=?',(g.owner_id,)).fetchone()
        from .drafts import reply_drafts
        from .ai_offers import ai_enabled,latest_for_lead
        lead=enrich(Lead(**json.loads(row['payload'])))
        drafts=reply_drafts(lead,profile['profile_services'],profile['portfolio'],dict(draft_project) if draft_project else None)
        ai_offer=latest_for_lead(db(),g.owner_id,lid,draft_project['id'] if draft_project else None)
        tags=db().execute('SELECT tag FROM lead_tags WHERE user_id=? AND lead_id=? ORDER BY tag',(g.owner_id,lid)).fetchall()
        return render_template('detail.html',row=row,lead=lead,activity=activity,reminder=reminder,
                               members=members,assignment=assignment,drafts=drafts,tags=tags,
                               draft_projects=draft_projects,draft_project=draft_project,
                               ai_offer=ai_offer,ai_configured=ai_enabled())

    @app.post('/app/leads/<int:lid>/ai-offers')
    def generate_ai_offers(lid):
        from .teams import EDIT_LEADS,require_role
        from .ai_offers import AIOfferError,generate_for_lead,groq_offer_drafts
        require_role(*EDIT_LEADS);row=user_lead(lid)
        project_id=request.form.get('project_id','')
        if project_id and not project_id.isdigit():abort(400)
        project=None
        if project_id:
            project=db().execute('''SELECT p.* FROM project_leads pl JOIN projects p
                ON p.id=pl.project_id AND p.user_id=pl.user_id
                WHERE pl.lead_id=? AND p.id=? AND p.workspace_id=?''',
                (lid,int(project_id),g.workspace['id'])).fetchone()
            if not project:abort(404)
        profile=db().execute('SELECT profile_services,portfolio FROM user_preferences WHERE user_id=?',(g.owner_id,)).fetchone()
        lead=enrich(Lead(**json.loads(row['payload'])))
        try:
            generate_for_lead(db(),g.workspace['id'],g.owner_id,g.user['id'],lid,lead,
                profile['profile_services'],profile['portfolio'],dict(project) if project else None,
                current_app.config.get('AI_OFFER_GENERATOR',groq_offer_drafts))
        except AIOfferError as exc:flash(str(exc))
        else:flash('AI-офферы готовы. Проверьте факты и условия перед отправкой.')
        target=url_for('detail',lid=lid,reply_project=project_id) if project_id else url_for('detail',lid=lid)
        return redirect(target+'#ai-offers')

    @app.route('/app/settings',methods=['GET','POST'])
    def settings():
        from .teams import EDIT_SETTINGS,require_role
        if request.method=='POST':
            require_role(*EDIT_SETTINGS)
            budget=request.form.get('min_budget','');selected=request.form.getlist('topics')
            if not budget.isdigit() or int(budget)>100000000 or not set(selected)<=set(TOPICS):abort(400)
            save_preferences(db(),g.owner_id,min_budget=int(budget),topics=selected,
                show_without_budget=bool(request.form.get('no_budget')),show_possible_needs=bool(request.form.get('possible')),
                monitoring_active=bool(request.form.get('active')),portfolio=request.form.get('portfolio','')[:3000])
            flash('Настройки применены к уведомлениям.');return redirect(url_for('settings'))
        prefs=db().execute('SELECT * FROM user_preferences WHERE user_id=?',(g.owner_id,)).fetchone()
        return render_template('settings.html',prefs=prefs,topics=TOPICS,topic_categories=TOPIC_CATEGORIES)

    @app.get('/app/sources')
    def sources():
        q=request.args.get('q','')[:100]
        rows=db().execute('''SELECT username,title,members,online,checked_at,last_message_at,check_reason
            FROM telegram_sources WHERE enabled=1 AND status='active' AND title ILIKE ? ESCAPE '\\'
            ORDER BY priority,online DESC LIMIT 100''',(like_pattern(q),)).fetchall()
        return render_template('sources.html',rows=rows)

    @app.get('/admin')
    def admin():
        if not is_platform_admin(g.user):abort(403)
        health=db().execute('SELECT * FROM health ORDER BY source').fetchall()
        users=db().execute('SELECT display_name,status,registered_at,last_seen_at FROM app_users ORDER BY id DESC LIMIT 100').fetchall()
        return render_template('admin.html',health=health,users=users)

    @app.errorhandler(400)
    @app.errorhandler(403)
    @app.errorhandler(404)
    @app.errorhandler(413)
    @app.errorhandler(429)
    @app.errorhandler(500)
    def error(exc):
        from werkzeug.exceptions import SecurityError
        if isinstance(exc,SecurityError):return 'Недопустимый адрес сайта.',400
        messages={400:'Проверьте введённые данные.',403:'У вас нет доступа к этой странице.',429:'Слишком много запросов. Подождите минуту и повторите.',
            404:'Страница или лид не найдены.',413:'Слишком большой запрос.',500:'Не удалось обработать запрос. Попробуйте позже.'}
        if request.path.startswith('/api/'):
            return jsonify({'error':messages.get(exc.code,'Ошибка запроса.'),'status':exc.code}),exc.code
        if request.path=='/tg/auth':
            return jsonify({'ok':False,'error':messages.get(exc.code,'Ошибка запроса.')}),exc.code
        if is_tg_path(request.path):
            return render_template('tg/error.html',message=messages.get(exc.code,'Ошибка запроса.')),exc.code
        return render_template('error.html',message=messages.get(exc.code,'Ошибка запроса.')),exc.code

    from .web_projects import register_projects
    register_projects(app,db)
    from .reminders import register_routes
    register_routes(app,db)
    from .web_telegram import register_telegram_routes
    register_telegram_routes(app,db)
    from .web_replies import register_reply_routes
    register_reply_routes(app,db)
    from .teams import register_team_routes
    register_team_routes(app,db)
    from .lead_actions import register_lead_actions
    register_lead_actions(app,db)
    from .integrations import register_integrations
    register_integrations(app,db)
    from .demand import register_demand
    register_demand(app,db)
    from .web_tg import register_tg_routes
    register_tg_routes(app,db)
    return app
